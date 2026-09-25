"""Unit and integration tests for NexusCareers aggregator pipeline."""

import io
import pytest
import openpyxl
from services.collectors import compute_dedup_hash, is_scam, is_tech_role
from services.ai_extractor import OpportunityRecord, fallback_classify_record
from services.excel_builder import build_excel_workbook, TIER_CONFIG
from services.telegram_notifier import generate_executive_caption
from api.trigger import authenticate_request


def test_dedup_hash():
    """Verify deduplication hash handles case and whitespace normalization."""
    hash1 = compute_dedup_hash("Anthropic", "Research Engineer - Alignment")
    hash2 = compute_dedup_hash("  anthropic ", "research engineer - alignment   ")
    hash3 = compute_dedup_hash("Anthropic", "Research Engineer - Systems")

    assert hash1 == hash2, "Identical normalized company and title must produce same hash"
    assert hash1 != hash3, "Different titles must produce different hashes"


def test_anti_scam_filter():
    """Verify anti-scam heuristic blacklist catches predatory posts."""
    assert is_scam("Python Developer", "Requires a refundable registration fee of $50") is True
    assert is_scam("AI Trainee", "Mandatory training deposit required before onboarding") is True
    assert is_scam("Software Engineer", "Unpaid test assignment for 3 weeks") is True
    assert is_scam("Backend Engineer", "Standard competitive salary and equity package") is False


def test_tech_role_filtering():
    """Verify tech role classifier matches technical positions and drops non-technical."""
    assert is_tech_role("Senior Backend Engineer") is True
    assert is_tech_role("Founding AI Researcher") is True
    assert is_tech_role("DevOps Specialist") is True
    assert is_tech_role("Software Engineering Intern") is True
    assert is_tech_role("Chief Marketing Officer") is False
    assert is_tech_role("Executive Assistant") is False


def test_opportunity_record_pydantic():
    """Verify Pydantic model validation with strict Literal enums."""
    record = OpportunityRecord(
        company_name="Perplexity",
        job_title="Member of Technical Staff - Search",
        seniority_tier="Senior / Staff / Lead (5+ YOE)",
        technical_domain="GenAI / LLM / Agentic Systems",
        min_exp_years=5,
        max_exp_years=8,
        location="San Francisco, CA",
        workplace_type="On-site",
        salary_package="$220k - $280k",
        platform="Ashby",
        apply_url="https://jobs.ashbyhq.com/perplexity/123",
        core_tech_stack=["Python", "PyTorch", "Kubernetes"],
        why_it_matters="Leading LLM-powered search engine scaling rapidly.",
    )
    assert record.seniority_tier == "Senior / Staff / Lead (5+ YOE)"
    assert record.technical_domain == "GenAI / LLM / Agentic Systems"


def test_fallback_classify_record():
    """Verify heuristic classification without LLM produces valid OpportunityRecord."""
    raw = {
        "company_name": "Glean",
        "job_title": "Software Engineer Intern - Summer 2026",
        "location": "Remote",
        "platform": "Ashby",
        "apply_url": "https://jobs.ashbyhq.com/glean/intern",
        "description": "Work on enterprise search using Python and React",
        "raw_salary": "$50/hr",
    }
    record = fallback_classify_record(raw)
    assert record.seniority_tier == "Internship"
    assert record.workplace_type == "Remote"
    assert "Python" in record.core_tech_stack


def test_excel_workbook_structure():
    """Verify Excel builder creates 4 styled sheets with proper tab names and formulas."""
    records = [
        OpportunityRecord(
            company_name="Anthropic",
            job_title="ML Research Intern",
            seniority_tier="Internship",
            technical_domain="Core ML / Deep Learning / Vision / NLP",
            min_exp_years=0,
            max_exp_years=1,
            location="San Francisco, CA",
            workplace_type="On-site",
            salary_package="$60/hr",
            platform="Ashby",
            apply_url="https://jobs.ashbyhq.com/anthropic/intern",
            posted_date="24 Sep 2026",
            core_tech_stack=["Python", "PyTorch"],
            why_it_matters="Frontier safety research.",
        ),
        OpportunityRecord(
            company_name="Stripe",
            job_title="Software Engineer I - Payments",
            seniority_tier="Fresher / 0-2 YOE",
            technical_domain="Backend & Distributed Systems",
            min_exp_years=0,
            max_exp_years=2,
            location="Seattle, WA",
            workplace_type="Hybrid",
            salary_package="$140k - $165k",
            platform="Greenhouse",
            apply_url="https://boards.greenhouse.io/stripe/123",
            posted_date="23 Sep 2026",
            core_tech_stack=["Ruby", "Go", "Distributed Systems"],
            why_it_matters="Global financial infrastructure.",
        ),
        OpportunityRecord(
            company_name="Cursor",
            job_title="Systems Engineer",
            seniority_tier="Mid-Level (2-5 YOE)",
            technical_domain="Full-Stack / Infrastructure / DevOps",
            min_exp_years=3,
            max_exp_years=5,
            location="San Francisco, CA",
            workplace_type="On-site",
            salary_package="$180k - $240k",
            platform="Ashby",
            apply_url="https://jobs.ashbyhq.com/cursor/456",
            posted_date="24 Sep 2026",
            core_tech_stack=["TypeScript", "Rust", "C++"],
            why_it_matters="Next-gen AI code editor.",
        ),
        OpportunityRecord(
            company_name="OpenAI",
            job_title="Staff Systems Engineer - Infrastructure",
            seniority_tier="Senior / Staff / Lead (5+ YOE)",
            technical_domain="Backend & Distributed Systems",
            min_exp_years=6,
            max_exp_years=12,
            location="San Francisco, CA",
            workplace_type="On-site",
            salary_package="$300k - $450k",
            platform="Greenhouse",
            apply_url="https://boards.greenhouse.io/openai/789",
            posted_date="22 Sep 2026",
            core_tech_stack=["Python", "Kubernetes", "Ray", "Triton"],
            why_it_matters="Frontier AI supercomputing.",
        ),
    ]

    buffer = build_excel_workbook(records)
    assert isinstance(buffer, io.BytesIO)
    assert buffer.getbuffer().nbytes > 0

    # Load and inspect with openpyxl
    wb = openpyxl.load_workbook(buffer)
    expected_sheet_names = [cfg["sheet_name"] for cfg in TIER_CONFIG]
    assert wb.sheetnames == expected_sheet_names

    # Check Internship sheet
    ws_intern = wb["🎓 Internships"]
    assert ws_intern.cell(row=1, column=1).value == "Company"
    assert ws_intern.cell(row=2, column=1).value == "Anthropic"
    # Check Posted Date in column 9
    assert ws_intern.cell(row=1, column=9).value == "Posted Date"
    assert ws_intern.cell(row=2, column=9).value == "24 Sep 2026"
    # Check hyperlink formula in column 11
    link_cell = ws_intern.cell(row=2, column=11).value
    assert '=HYPERLINK("https://jobs.ashbyhq.com/anthropic/intern", "Apply Direct ↗")' in link_cell


def test_telegram_caption_generation():
    """Verify executive summary telegram caption generation."""
    records = [
        OpportunityRecord(
            company_name="Anthropic",
            job_title="AI Research Intern",
            seniority_tier="Internship",
            technical_domain="GenAI / LLM / Agentic Systems",
            location="San Francisco",
            workplace_type="On-site",
            salary_package="$60/hr",
            platform="Ashby",
            apply_url="https://...",
            core_tech_stack=["PyTorch"],
            why_it_matters="Leading frontier AI.",
        )
    ]
    caption = generate_executive_caption(records)
    assert "TechyUpdates" in caption
    assert "🎓 *Internships & College Grads:* 1" in caption
    assert "Anthropic" in caption


def test_authenticate_request():
    """Verify Bearer token matching logic."""
    import os
    os.environ["CRON_SECRET"] = "my_secret_token_123"

    assert authenticate_request({"Authorization": "Bearer my_secret_token_123"}) is True
    assert authenticate_request({"Authorization": "Bearer wrong_token"}) is False
    assert authenticate_request({}) is False

    del os.environ["CRON_SECRET"]
    # In dev mode without CRON_SECRET set, it permits
    assert authenticate_request({}) is True


def test_liveness_closure_markers():
    """Verify that expired and closed markers are properly identified."""
    from services.collectors.liveness_verifier import CLOSURE_MARKERS

    sample_html = "<html><body><h1>Sorry, this job is closed.</h1><p>We are no longer accepting applications.</p></body></html>"
    found = any(m in sample_html.lower() for m in CLOSURE_MARKERS)
    assert found is True

    active_html = "<html><body><h1>Apply for Software Engineer</h1><form action='/submit'></form></body></html>"
    active_found = any(m in active_html.lower() for m in CLOSURE_MARKERS)
    assert active_found is False


def test_health_check_endpoint():
    """Verify health check logic returns 200 without running pipeline."""
    from api.trigger import handler
    from unittest.mock import MagicMock

    h = handler.__new__(handler)
    h.path = "/api/health"
    h.headers = {}
    h._send_response_json = MagicMock()

    h.do_GET()

    h._send_response_json.assert_called_once()
    args, _ = h._send_response_json.call_args
    assert args[0] == 200
    assert args[1]["status"] == "healthy"
    assert args[1]["service"] == "techyupdates-careers-engine"


